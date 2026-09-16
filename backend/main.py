from fastapi import FastAPI, Query, HTTPException, Depends, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import select, func, distinct, delete, text
from database import async_session, SocialObject, User, Builder
from auth import verify_password, hash_password, create_access_token, get_current_admin
from pydantic import BaseModel
from typing import Optional, List
from datetime import date, datetime
import json
import csv
import io
import os
import time
import openpyxl
import httpx

app = FastAPI(title="Тульская область - Портал ОКС")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ObjectResponse(BaseModel):
    id: int
    num: Optional[int] = None
    grbs: Optional[str] = None
    oks_name: Optional[str] = None
    construction_stage: Optional[str] = None
    address: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    industry: Optional[str] = None
    status: Optional[str] = None
    ownership: Optional[str] = None
    amo: Optional[str] = None
    customer: Optional[str] = None
    np_gp_name: Optional[str] = None
    fp_name: Optional[str] = None
    project_code: Optional[str] = None
    total_area_m2: Optional[float] = None
    capacity: Optional[str] = None
    expertise: Optional[str] = None
    year_start: Optional[int] = None
    year_end: Optional[int] = None
    construction_period: Optional[str] = None
    land_transfer_date: Optional[date] = None
    building_permit_date: Optional[date] = None
    contract_date: Optional[date] = None
    contract_period: Optional[str] = None
    contractor: Optional[str] = None
    construction_readiness: Optional[float] = None
    equipment_installation_date: Optional[date] = None
    hydro_test_date: Optional[date] = None
    zos_date: Optional[date] = None
    zos_number: Optional[str] = None
    act_input_date: Optional[date] = None
    act_input_number: Optional[str] = None
    year_commissioned: Optional[int] = None
    photo_before: Optional[str] = None
    photo_after: Optional[str] = None

    class Config:
        from_attributes = True

class ObjectListItem(BaseModel):
    id: int
    num: Optional[int] = None
    oks_name: Optional[str] = None
    construction_stage: Optional[str] = None
    address: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    industry: Optional[str] = None
    status: Optional[str] = None
    ownership: Optional[str] = None
    amo: Optional[str] = None
    customer: Optional[str] = None
    np_gp_name: Optional[str] = None
    fp_name: Optional[str] = None
    project_code: Optional[str] = None
    total_area_m2: Optional[float] = None
    capacity: Optional[str] = None
    expertise: Optional[str] = None
    year_start: Optional[int] = None
    year_end: Optional[int] = None
    construction_period: Optional[str] = None
    contractor: Optional[str] = None
    construction_readiness: Optional[float] = None
    year_commissioned: Optional[int] = None
    grbs: Optional[str] = None
    land_transfer_date: Optional[date] = None
    building_permit_date: Optional[date] = None
    contract_date: Optional[date] = None
    contract_period: Optional[str] = None
    equipment_installation_date: Optional[date] = None
    hydro_test_date: Optional[date] = None
    zos_date: Optional[date] = None
    zos_number: Optional[str] = None
    act_input_date: Optional[date] = None
    act_input_number: Optional[str] = None

    class Config:
        from_attributes = True

class YearStats(BaseModel):
    year: int
    count: int

class BuilderResponse(BaseModel):
    id: int
    municipality: Optional[str] = None
    category: Optional[str] = None
    title: str
    address: Optional[str] = None
    longitude: Optional[float] = None
    latitude: Optional[float] = None

    class Config:
        from_attributes = True

class MunicipalityColorsResponse(BaseModel):
    municipality: str
    population_total: Optional[int] = None
    schools_per_1000: Optional[float] = None
    color_school: Optional[str] = None
    hospitals_count: Optional[int] = None
    people_per_hospital: Optional[float] = None
    color_hospital: Optional[str] = None
    sports_count: Optional[int] = None
    people_per_sport: Optional[float] = None
    color_sport: Optional[str] = None

    class Config:
        from_attributes = True

class LoginRequest(BaseModel):
    username: str
    password: str

class LoginResponse(BaseModel):
    token: str
    username: str

class UserResponse(BaseModel):
    id: int
    username: str
    created_at: Optional[datetime] = None

    class Config:
        from_attributes = True

class UserCreateRequest(BaseModel):
    username: str
    password: str

class PasswordChangeRequest(BaseModel):
    old_password: str
    new_password: str

class BulkImportResult(BaseModel):
    inserted: int
    updated: int
    errors: List[str]

@app.get("/api/objects", response_model=List[ObjectListItem])
async def get_objects(
    year_start: Optional[List[int]] = Query(None, description="Год начала строительства"),
    year_end: Optional[List[int]] = Query(None, description="Год окончания"),
    status: Optional[str] = Query(None, description="Статус объекта"),
    industry: Optional[str] = Query(None, description="Отрасль"),
    construction_stage: Optional[str] = Query(None, description="Этап строительства"),
    search: Optional[str] = Query(None, description="Поиск по наименованию/адресу"),
    lat_min: Optional[float] = Query(None),
    lat_max: Optional[float] = Query(None),
    lon_min: Optional[float] = Query(None),
    lon_max: Optional[float] = Query(None),
):

    list_columns = [
        SocialObject.id, SocialObject.num, SocialObject.oks_name,
        SocialObject.construction_stage, SocialObject.address,
        SocialObject.latitude, SocialObject.longitude, SocialObject.industry,
        SocialObject.status, SocialObject.ownership, SocialObject.amo,
        SocialObject.customer, SocialObject.np_gp_name, SocialObject.fp_name,
        SocialObject.project_code, SocialObject.total_area_m2, SocialObject.capacity,
        SocialObject.expertise, SocialObject.year_start, SocialObject.year_end,
        SocialObject.construction_period, SocialObject.contractor,
        SocialObject.construction_readiness, SocialObject.year_commissioned,
        SocialObject.grbs, SocialObject.land_transfer_date,
        SocialObject.building_permit_date, SocialObject.contract_date,
        SocialObject.contract_period, SocialObject.equipment_installation_date,
        SocialObject.hydro_test_date, SocialObject.zos_date, SocialObject.zos_number,
        SocialObject.act_input_date, SocialObject.act_input_number,
    ]
    async with async_session() as session:
        query = select(*list_columns).where(SocialObject.latitude.isnot(None), SocialObject.longitude.isnot(None))

        if year_start:
            query = query.where(SocialObject.year_start.in_(year_start))
        if year_end:
            query = query.where(SocialObject.year_end.in_(year_end))
        if status:
            query = query.where(SocialObject.status.ilike(f"%{status}%"))
        if industry:
            query = query.where(SocialObject.industry == industry)
        if construction_stage:
            if construction_stage == 'Строится':

                query = query.where(SocialObject.construction_stage.is_(None))
            else:
                query = query.where(SocialObject.construction_stage.ilike(f"%{construction_stage}%"))
        if search:
            query = query.where(
                (SocialObject.oks_name.ilike(f"%{search}%")) |
                (SocialObject.address.ilike(f"%{search}%"))
            )
        if lat_min is not None:
            query = query.where(SocialObject.latitude >= lat_min)
        if lat_max is not None:
            query = query.where(SocialObject.latitude <= lat_max)
        if lon_min is not None:
            query = query.where(SocialObject.longitude >= lon_min)
        if lon_max is not None:
            query = query.where(SocialObject.longitude <= lon_max)

        result = await session.execute(query)
        return [row._mapping for row in result.all()]

@app.get("/api/objects/{obj_id}", response_model=ObjectResponse)
async def get_object(obj_id: int):
    async with async_session() as session:
        query = select(SocialObject).where(SocialObject.id == obj_id)
        result = await session.execute(query)
        obj = result.scalar_one_or_none()
        if not obj:
            raise HTTPException(status_code=404, detail="Объект не найден")
        return obj

@app.get("/api/objects/{obj_id}/photos")
async def get_object_photos(obj_id: int):
    async with async_session() as session:
        result = await session.execute(
            select(SocialObject.photo_before, SocialObject.photo_after)
            .where(SocialObject.id == obj_id)
        )
        row = result.first()
        if not row:
            raise HTTPException(status_code=404, detail="Объект не найден")
        return {"photo_before": row[0], "photo_after": row[1]}

@app.get("/api/years", response_model=List[YearStats])
async def get_years():
    async with async_session() as session:
        query = (
            select(SocialObject.year_start, func.count(SocialObject.id))
            .where(SocialObject.year_start.isnot(None))
            .group_by(SocialObject.year_start)
            .order_by(SocialObject.year_start)
        )
        result = await session.execute(query)
        return [{"year": row[0], "count": row[1]} for row in result.all()]

@app.get("/api/stats")
async def get_stats():
    async with async_session() as session:
        total = await session.execute(select(func.count(SocialObject.id)))
        total_count = total.scalar()

        completed = await session.execute(
            select(func.count(SocialObject.id)).where(
                SocialObject.construction_stage.ilike('%введен%')
            )
        )
        completed_count = completed.scalar()

        active = await session.execute(
            select(func.count(SocialObject.id)).where(
                SocialObject.year_end.is_(None)
            )
        )
        active_count = active.scalar()

        stages = await session.execute(
            select(SocialObject.construction_stage, func.count(SocialObject.id))
            .group_by(SocialObject.construction_stage)
        )
        stage_stats = {row[0] or "Не указан": row[1] for row in stages.all()}

        industries = await session.execute(
            select(SocialObject.industry, func.count(SocialObject.id))
            .group_by(SocialObject.industry)
        )
        industry_stats = {row[0] or "Не указана": row[1] for row in industries.all()}

        return {
            "total": total_count,
            "completed": completed_count,
            "active": active_count,
            "by_stage": stage_stats,
            "by_industry": industry_stats,
        }

@app.get("/api/industries")
async def get_industries():
    async with async_session() as session:
        query = (
            select(SocialObject.industry)
            .where(SocialObject.industry.isnot(None))
            .distinct()
            .order_by(SocialObject.industry)
        )
        result = await session.execute(query)
        return [row[0] for row in result.all()]

@app.get("/api/stages")
async def get_stages():
    async with async_session() as session:
        query = (
            select(SocialObject.construction_stage)
            .where(SocialObject.construction_stage.isnot(None))
            .distinct()
            .order_by(SocialObject.construction_stage)
        )
        result = await session.execute(query)
        return [row[0] for row in result.all()]

@app.get("/api/builders", response_model=List[BuilderResponse])
async def get_builders(
    category: Optional[str] = Query(None, description="Категория объекта"),
):
    async with async_session() as session:
        query = select(Builder).where(
            Builder.latitude.isnot(None), Builder.longitude.isnot(None)
        )
        if category:
            query = query.where(Builder.category == category)
        query = query.order_by(Builder.title)
        result = await session.execute(query)
        return result.scalars().all()

@app.get("/api/municipalities-colors", response_model=List[MunicipalityColorsResponse])
async def get_municipalities_colors():
    async with async_session() as session:
        result = await session.execute(text("SELECT * FROM map_build.municipalities_colors"))
        return [dict(row) for row in result.mappings().all()]

YANDEX_ISOLINE_API_KEY = os.getenv("YANDEX_ISOLINE_API_KEY", "34bcc13c-1716-45bb-b14d-78d9822422ea")
YANDEX_ISOLINE_URL = "https://isoline.api.maps.yandex.ru/v1/{mode}"

ISOLINE_MODES = {"driving", "walking", "transit", "bicycle", "scooter"}
ISOLINE_JAMS_TYPES = {"historical", "realtime", "disabled"}

_ISOLINE_CACHE_TTL = 3600
_isoline_cache: dict = {}

@app.get("/api/isochrone")
async def get_isochrone(
    lat: float = Query(..., description="Широта точки (WGS84)"),
    lon: float = Query(..., description="Долгота точки (WGS84)"),
    duration: int = Query(900, description="Время в пути, секунды"),
    mode: str = Query("walking", description="Способ перемещения"),
    jams_type: Optional[str] = Query(None, description="Тип пробок (только для driving)"),
):
    if mode not in ISOLINE_MODES:
        raise HTTPException(status_code=400, detail=f"Недопустимый mode. Допустимые значения: {', '.join(sorted(ISOLINE_MODES))}")

    max_duration = 3600 if mode == "driving" else 2700
    if not (0 < duration <= max_duration):
        raise HTTPException(status_code=400, detail=f"duration должен быть от 1 до {max_duration} секунд для mode={mode}")

    if jams_type is not None:
        if mode != "driving":
            raise HTTPException(status_code=400, detail="jams_type применим только для mode=driving")
        if jams_type not in ISOLINE_JAMS_TYPES:
            raise HTTPException(status_code=400, detail=f"Недопустимый jams_type. Допустимые значения: {', '.join(sorted(ISOLINE_JAMS_TYPES))}")

    cache_key = (mode, round(lat, 5), round(lon, 5), duration, jams_type)
    cached = _isoline_cache.get(cache_key)
    if cached and (time.monotonic() - cached[0]) < _ISOLINE_CACHE_TTL:
        return cached[1]

    params = {
        "apikey": YANDEX_ISOLINE_API_KEY,
        "ll": f"{lon},{lat}",
        "duration": duration,
    }
    if jams_type:
        params["jams_type"] = jams_type

    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.get(YANDEX_ISOLINE_URL.format(mode=mode), params=params)
    except httpx.HTTPError:
        raise HTTPException(status_code=502, detail="Не удалось связаться с сервисом изохрон")

    if resp.status_code != 200:
        raise HTTPException(status_code=502, detail=f"Сервис изохрон вернул ошибку ({resp.status_code})")

    data = resp.json()
    _isoline_cache[cache_key] = (time.monotonic(), data)
    return data

@app.post("/api/auth/login", response_model=LoginResponse)
async def login(req: LoginRequest):
    async with async_session() as session:
        result = await session.execute(select(User).where(User.username == req.username))
        user = result.scalar_one_or_none()

        token = create_access_token({"sub": user.username})
        return LoginResponse(token=token, username=user.username)

@app.get("/api/auth/me", response_model=UserResponse)
async def get_me(admin: User = Depends(get_current_admin)):
    return admin

@app.get("/api/admin/users", response_model=List[UserResponse])
async def list_users(admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        result = await session.execute(select(User).order_by(User.id))
        return result.scalars().all()

@app.post("/api/admin/users", response_model=UserResponse)
async def create_user(req: UserCreateRequest, admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        existing = await session.execute(select(User).where(User.username == req.username))
        if existing.scalar_one_or_none():
            raise HTTPException(status_code=400, detail="Пользователь уже существует")
        user = User(username=req.username, password=hash_password(req.password))
        session.add(user)
        await session.commit()
        await session.refresh(user)
        return user

@app.delete("/api/admin/users/{user_id}")
async def delete_user(user_id: int, admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        result = await session.execute(select(User).where(User.id == user_id))
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(status_code=404, detail="Пользователь не найден")
        if user.username == "admin":
            raise HTTPException(status_code=400, detail="Нельзя удалить главного администратора")
        await session.delete(user)
        await session.commit()
        return {"detail": "Пользователь удалён"}

@app.post("/api/admin/users/{user_id}/password")
async def change_password(user_id: int, req: PasswordChangeRequest, admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        result = await session.execute(select(User).where(User.id == user_id))
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(status_code=404, detail="Пользователь не найден")
        if not verify_password(req.old_password, user.password):
            raise HTTPException(status_code=400, detail="Неверный текущий пароль")
        user.password = hash_password(req.new_password)
        await session.commit()
        return {"detail": "Пароль обновлён"}

OBJECT_COLUMNS = [
    "num", "grbs", "oks_name", "construction_stage", "address",
    "latitude", "longitude", "industry", "status", "ownership",
    "amo", "customer", "np_gp_name", "fp_name", "project_code",
    "total_area_m2", "capacity", "expertise", "year_start", "year_end",
    "construction_period", "land_transfer_date", "building_permit_date",
    "contract_date", "contract_period", "contractor", "construction_readiness",
    "equipment_installation_date", "hydro_test_date", "zos_date", "zos_number",
    "act_input_date", "act_input_number", "year_commissioned", "photo_before", "photo_after",
]

DATE_FIELDS = {
    "land_transfer_date", "building_permit_date", "contract_date",
    "equipment_installation_date", "hydro_test_date", "zos_date",
    "act_input_date",
}

FLOAT_FIELDS = {
    "latitude", "longitude", "total_area_m2", "construction_readiness",
}

INT_FIELDS = {
    "num", "year_start", "year_end", "year_commissioned",
}

def parse_cell(value, field: str):
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if field in DATE_FIELDS:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, str):
            try:
                return date.fromisoformat(value.strip())
            except ValueError:
                return None
        return None
    if field in FLOAT_FIELDS:
        try:
            return float(value)
        except (ValueError, TypeError):
            return None
    if field in INT_FIELDS:
        try:
            return int(float(value))
        except (ValueError, TypeError):
            return None
    return str(value).strip() if isinstance(value, str) else value

def parse_rows_from_csv(content: str) -> List[dict]:
    reader = csv.DictReader(io.StringIO(content))
    rows = []
    for row in reader:
        normalized = {}
        for col in OBJECT_COLUMNS:
            key = col if col in row else col.lower()
            normalized[col] = parse_cell(row.get(key), col)
        rows.append(normalized)
    return rows

def parse_rows_from_xlsx(file_bytes: bytes) -> List[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers_raw = next(rows_iter, None)
    if not headers_raw:
        return []

    headers = [str(h).strip().lower() if h else "" for h in headers_raw]
    col_map = {}
    for i, h in enumerate(headers):
        if h in OBJECT_COLUMNS:
            col_map[h] = i
        elif h.replace(" ", "_") in OBJECT_COLUMNS:
            col_map[h.replace(" ", "_")] = i

    rows = []
    for row_cells in rows_iter:
        normalized = {}
        for col in OBJECT_COLUMNS:
            idx = col_map.get(col)
            val = row_cells[idx] if idx is not None and idx < len(row_cells) else None
            normalized[col] = parse_cell(val, col)
        rows.append(normalized)
    wb.close()
    return rows

@app.post("/api/admin/import", response_model=BulkImportResult)
async def import_objects(file: UploadFile = File(...), admin: User = Depends(get_current_admin)):
    if not file.filename.endswith(('.csv', '.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Поддерживаются CSV и Excel (.xlsx) файлы")

    file_bytes = await file.read()

    if file.filename.endswith('.csv'):
        content = file_bytes.decode("utf-8-sig")
        rows = parse_rows_from_csv(content)
    else:
        rows = parse_rows_from_xlsx(file_bytes)

    if not rows:
        raise HTTPException(status_code=400, detail="Файл пуст или не содержит данных")

    inserted = 0
    updated = 0
    errors = []

    async with async_session() as session:
        for i, row_data in enumerate(rows, start=1):
            try:
                obj_id = row_data.pop("num", None)

                existing = None
                if obj_id is not None:
                    result = await session.execute(
                        select(SocialObject).where(SocialObject.num == obj_id)
                    )
                    existing = result.scalar_one_or_none()

                if existing:
                    for field, value in row_data.items():
                        setattr(existing, field, value)
                    updated += 1
                else:
                    row_data["num"] = obj_id
                    obj = SocialObject(**row_data)
                    session.add(obj)
                    inserted += 1
            except Exception as e:
                errors.append(f"Строка {i}: {str(e)}")

        await session.commit()

    return BulkImportResult(inserted=inserted, updated=updated, errors=errors)

@app.get("/api/admin/export")
async def export_objects(format: str = "csv", admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        result = await session.execute(select(SocialObject).order_by(SocialObject.id))
        objects = result.scalars().all()

    if format == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Объекты"
        headers = ["id"] + OBJECT_COLUMNS
        ws.append(headers)
        for obj in objects:
            row = [getattr(obj, col, None) for col in ["id"] + OBJECT_COLUMNS]
            ws.append(row)

        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        buf.seek(0)
        from starlette.responses import StreamingResponse
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=objects.xlsx"},
        )

    output = io.StringIO()
    writer = csv.writer(output)
    headers = ["id"] + OBJECT_COLUMNS
    writer.writerow(headers)
    for obj in objects:
        row = [getattr(obj, col, None) for col in ["id"] + OBJECT_COLUMNS]
        writer.writerow(row)

    from starlette.responses import Response
    return Response(
        content=output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=objects.csv"},
    )

@app.get("/api/admin/stats")
async def admin_stats(admin: User = Depends(get_current_admin)):
    async with async_session() as session:
        total_objects = (await session.execute(select(func.count(SocialObject.id)))).scalar()
        total_users = (await session.execute(select(func.count(User.id)))).scalar()

        stages = await session.execute(
            select(SocialObject.construction_stage, func.count(SocialObject.id))
            .group_by(SocialObject.construction_stage)
        )
        stage_stats = {row[0] or "Не указан": row[1] for row in stages.all()}

        industries = await session.execute(
            select(SocialObject.industry, func.count(SocialObject.id))
            .group_by(SocialObject.industry)
        )
        industry_stats = {row[0] or "Не указана": row[1] for row in industries.all()}

        readiness = await session.execute(
            select(func.avg(SocialObject.construction_readiness))
            .where(SocialObject.construction_readiness.isnot(None))
            .where(SocialObject.construction_readiness > 0)
        )
        avg_readiness = readiness.scalar()

        completed = await session.execute(
            select(func.count(SocialObject.id))
            .where(SocialObject.year_commissioned.isnot(None))
        )
        completed_count = completed.scalar()

        return {
            "total_objects": total_objects,
            "total_users": total_users,
            "avg_readiness": round(avg_readiness, 1) if avg_readiness else 0,
            "completed_count": completed_count,
            "active_count": total_objects - completed_count,
            "by_stage": stage_stats,
            "by_industry": industry_stats,
        }
